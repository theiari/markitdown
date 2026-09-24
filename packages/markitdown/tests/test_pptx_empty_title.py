#!/usr/bin/env python3 -m pytest
"""A slide whose title placeholder carries no text must not get a bare heading."""

import io

import pptx

from markitdown import MarkItDown, StreamInfo

BODY_TEXT = "Some body text on the slide."


def _build_pptx(title: str | None) -> io.BytesIO:
    """Build a one-slide "Title and Content" deck, optionally filling the title.

    Passing None leaves the title placeholder as PowerPoint creates it: present
    on the slide, showing "Click to add title", and carrying no text.
    """
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    if title is not None:
        slide.shapes.title.text = title
    slide.placeholders[1].text = BODY_TEXT
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


def _convert(title: str | None) -> str:
    return (
        MarkItDown()
        .convert_stream(_build_pptx(title), stream_info=StreamInfo(extension=".pptx"))
        .markdown
    )


def _heading_lines(markdown: str) -> list[str]:
    return [line for line in markdown.splitlines() if line.startswith("#")]


def test_untouched_title_placeholder_produces_no_heading() -> None:
    markdown = _convert(None)

    assert BODY_TEXT in markdown
    assert _heading_lines(markdown) == []


def test_empty_title_produces_no_heading() -> None:
    markdown = _convert("")

    assert BODY_TEXT in markdown
    assert _heading_lines(markdown) == []


def test_whitespace_only_title_produces_no_heading() -> None:
    markdown = _convert("   ")

    assert BODY_TEXT in markdown
    assert _heading_lines(markdown) == []


def test_title_with_text_is_still_emitted() -> None:
    markdown = _convert("Quarterly Results")

    assert "# Quarterly Results" in markdown
    assert BODY_TEXT in markdown
