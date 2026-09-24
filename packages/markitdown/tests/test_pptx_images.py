"""PPTX image hooks share native slide traversal and HTML rendering."""

import base64
import inspect
import io
from pathlib import Path
from typing import Any, BinaryIO, Callable, Optional
from unittest.mock import Mock

from bs4 import BeautifulSoup
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches
import pytest

from markitdown import (
    FileConversionException,
    MarkItDown,
    MissingDependencyException,
    StreamInfo,
)
from markitdown.converters import HtmlConverter, PptxConverter
from markitdown.converters import _pptx_converter


_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAIAAACQd1PeAAAADElEQVR4nGP4z8AAAAM"
    "BAQDJ/pLvAAAAAElFTkSuQmCC"
)
_GIF = base64.b64decode("R0lGODlhAQABAIAAAAAAAP///yH5BAEAAAAALAAAAAABAAEAAAIBTAA7")
_INFO = StreamInfo(extension=".pptx")
_FILES = Path(__file__).parent / "test_files"


def _presentation(
    images: tuple[bytes, ...] = (_PNG,), *, native_content: bool = False
) -> io.BytesIO:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Deck title"
    slide.shapes.title.top = 0
    # Insert in reverse order to distinguish shape order from reading order.
    for index in reversed(range(len(images))):
        image = slide.shapes.add_picture(
            io.BytesIO(images[index]), 0, Inches(index + 1), width=Inches(0.5)
        )
        image.name = f"Picture {index + 1}"
        image._element._nvXxPr.cNvPr.set("descr", f"Alt [{index + 1}]\r\n& text")
    if native_content:
        table = slide.shapes.add_table(2, 2, 0, Inches(4), Inches(4), Inches(1)).table
        for cell, value in zip(
            (table.cell(0, 0), table.cell(0, 1), table.cell(1, 0), table.cell(1, 1)),
            ("Item", "Details", "A", "B & <literal>"),
        ):
            cell.text = value
        chart_data = CategoryChartData()
        chart_data.categories = ["Q1", "Q2"]
        chart_data.add_series("Revenue", (10, 20))
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            0,
            Inches(5),
            Inches(4),
            Inches(1),
            chart_data,
        ).chart
        chart.has_title = True
        chart.chart_title.text_frame.text = "Sales"
        slide.notes_slide.notes_text_frame.text = "Speaker notes"
        next_slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        next_slide.shapes.title.text = ""
        next_slide.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = "Closing"
        next_slide.notes_slide.notes_text_frame.text = " \n "
    stream = io.BytesIO()
    presentation.save(stream)
    stream.seek(0)
    return stream


class _ImageConverter(PptxConverter):
    def __init__(self, render: Callable[..., Optional[str]]):
        super().__init__()
        self.render = render

    def _image_to_html(
        self, image_stream: BinaryIO, stream_info: StreamInfo, **kwargs: Any
    ) -> Optional[str]:
        return self.render(image_stream, stream_info, **kwargs)


def test_hook_does_not_change_public_signature() -> None:
    assert list(inspect.signature(PptxConverter.convert).parameters) == [
        "self",
        "file_stream",
        "stream_info",
        "kwargs",
    ]
    assert list(inspect.signature(PptxConverter.__init__).parameters) == ["self"]
    assert _ImageConverter.convert is PptxConverter.convert


@pytest.mark.parametrize("via_dispatcher", [False, True])
def test_inherited_hook_receives_real_images_and_options_in_reading_order(
    via_dispatcher: bool,
) -> None:
    seen = []
    streams = []
    service = object()

    def render(stream: BinaryIO, info: StreamInfo, **kwargs: Any) -> str:
        assert stream.tell() == 0 and stream.seekable()
        data = stream.read()
        stream.seek(0)
        assert stream.read() == data
        seen.append((data, info, kwargs))
        streams.append(stream)
        return f"<strong>Image {len(seen)}</strong>"

    class InheritedImages(_ImageConverter):
        pass

    converter = InheritedImages(render)
    source = _presentation((_GIF, _PNG, _GIF))
    original = source.getvalue()
    info = StreamInfo(
        extension=".pptx",
        filename="deck.pptx",
        url="https://example.test/deck.pptx",
        local_path="/deck.pptx",
    )
    options: dict[str, Any] = {"ocr_service": service, "escape_underscores": False}
    if via_dispatcher:
        md = MarkItDown()
        md.register_converter(converter, priority=-1)
        result = md.convert_stream(source, stream_info=info, **options)
    else:
        source.seek(7)
        result = converter.convert(source, info, **options)

    assert [entry[0] for entry in seen] == [_GIF, _PNG, _GIF]
    assert [entry[1] for entry in seen] == [
        StreamInfo(
            mimetype=f"image/{ext}", extension=f".{ext}", filename=f"image.{ext}"
        )
        for ext in ("gif", "png", "gif")
    ]
    assert all(entry[2]["ocr_service"] is service for entry in seen)
    assert all(entry[2]["escape_underscores"] is False for entry in seen)
    assert result.markdown == (
        "<!-- Slide number: 1 -->\n# Deck title\n"
        "\n**Image 1**\n\n**Image 2**\n\n**Image 3**"
    )
    assert result.title is None
    assert all(stream.closed for stream in streams)
    assert not source.closed and source.getvalue() == original


@pytest.mark.parametrize(
    ("fallback", "keep_data_uris"), [(None, False), ("", True), (" \r\n\t", False)]
)
def test_declining_hook_preserves_native_output_and_metadata(
    fallback: Optional[str], keep_data_uris: bool
) -> None:
    expected = PptxConverter().convert(
        _presentation(native_content=True), _INFO, keep_data_uris=keep_data_uris
    )
    render = Mock(return_value=fallback)
    actual = _ImageConverter(render).convert(
        _presentation(native_content=True), _INFO, keep_data_uris=keep_data_uris
    )

    assert actual.markdown == expected.markdown
    assert actual.title == expected.title
    render.assert_called_once()


@pytest.mark.parametrize("keep_data_uris", [False, True])
def test_declining_hook_preserves_real_presentation(keep_data_uris: bool) -> None:
    data = (_FILES / "test.pptx").read_bytes()
    expected = PptxConverter().convert(
        io.BytesIO(data), _INFO, keep_data_uris=keep_data_uris
    )
    render = Mock(return_value=None)

    actual = _ImageConverter(render).convert(
        io.BytesIO(data), _INFO, keep_data_uris=keep_data_uris
    )

    assert actual.markdown == expected.markdown
    assert actual.title == expected.title
    assert render.call_count > 0


def test_native_converter_does_not_add_image_html_processing(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class NativeImages(PptxConverter):
        pass

    for converter in (PptxConverter(), NativeImages()):
        convert_html = Mock(side_effect=AssertionError("unexpected HTML"))
        monkeypatch.setattr(converter._html_converter, "convert_string", convert_html)
        assert (
            "![Alt 1 & text](Picture1.jpg)"
            in converter.convert(_presentation(), _INFO).markdown
        )
        convert_html.assert_not_called()


def test_override_can_delegate_to_super() -> None:
    class NativeImages(PptxConverter):
        def _image_to_html(self, image_stream, stream_info, **kwargs):
            return super()._image_to_html(image_stream, stream_info, **kwargs)

    assert (
        NativeImages().convert(_presentation(), _INFO).markdown
        == PptxConverter().convert(_presentation(), _INFO).markdown
    )


def test_custom_fragment_uses_html_options_without_reprocessing_generated_images(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    fragment = (
        "<h2>Image heading</h2><p>A_B &amp; &lt;literal&gt;<br>second</p>"
        '<img src="generated.png" alt="generated"><script>discard</script>'
    )
    render = Mock(return_value=fragment)
    converter = _ImageConverter(render)
    convert_html = Mock(wraps=converter._html_converter.convert_string)
    monkeypatch.setattr(converter._html_converter, "convert_string", convert_html)
    options: dict[str, Any] = {
        "escape_underscores": False,
        "heading_style": "underlined",
    }

    result = converter.convert(_presentation(), _INFO, **options)

    convert_html.assert_called_once_with(
        str(BeautifulSoup(fragment, "html.parser")), **options
    )
    expected = HtmlConverter().convert_string(fragment, **options).markdown
    assert result.markdown == "<!-- Slide number: 1 -->\n# Deck title\n\n" + expected
    render.assert_called_once()


def test_replacements_keep_native_table_chart_notes_and_slide_placement() -> None:
    render = Mock(side_effect=["<p>first</p>", None, "<p>third</p>"])
    result = _ImageConverter(render).convert(
        _presentation((_PNG, _GIF, _PNG), native_content=True), _INFO
    )
    markdown = result.markdown
    ordered = [
        "# Deck title",
        "first",
        "![Alt 2 & text](Picture2.jpg)",
        "third",
        "| Item | Details |",
        "| A | B & <literal> |",
        "### Chart: Sales",
        "| Q1 | 10.0 |",
        "### Notes:\nSpeaker notes",
        "<!-- Slide number: 2 -->",
        "Closing",
    ]
    assert [markdown.index(text) for text in ordered] == sorted(
        markdown.index(text) for text in ordered
    )
    assert markdown.count("### Notes:") == 1
    assert "\n# \n" not in markdown


def test_group_images_follow_native_negative_and_zero_coordinate_order() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_picture(io.BytesIO(_PNG), 0, 0, width=Inches(0.5))
    group.shapes.add_picture(io.BytesIO(_GIF), 0, -Inches(1), width=Inches(0.5))
    stream = io.BytesIO()
    presentation.save(stream)
    stream.seek(0)
    seen = []

    def render(image_stream, info, **kwargs):
        seen.append(image_stream.read())
        return f"<p>image{len(seen)}</p>"

    result = _ImageConverter(render).convert(stream, _INFO)

    assert seen == [_GIF, _PNG]
    assert result.markdown == "<!-- Slide number: 1 -->\n\nimage1\n\nimage2"


def test_picture_placeholder_invokes_hook() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[8])
    placeholder = next(
        shape for shape in slide.placeholders if hasattr(shape, "insert_picture")
    )
    placeholder.insert_picture(io.BytesIO(_PNG))
    stream = io.BytesIO()
    presentation.save(stream)
    stream.seek(0)
    render = Mock(return_value="<p>placeholder</p>")

    result = _ImageConverter(render).convert(stream, _INFO)

    assert result.markdown == "<!-- Slide number: 1 -->\n\nplaceholder"
    render.assert_called_once()


def test_svg_without_raster_fallback_reaches_image_hook() -> None:
    seen = []

    def render(stream, info, **kwargs):
        seen.append((stream.read(), info))
        return "<p>SVG content</p>"

    with (_FILES / "test_svg_no_fallback.pptx").open("rb") as stream:
        result = _ImageConverter(render).convert(stream, _INFO)

    assert len(seen) == 1 and b"<svg" in seen[0][0]
    assert seen[0][1].mimetype == "image/svg+xml"
    assert seen[0][1].extension == ".svg"
    assert seen[0][1].url is None
    assert "SVG content" in result.markdown


def test_svg_picture_placeholder_uses_native_resolution_and_image_hook() -> None:
    presentation = Presentation(str(_FILES / "test_svg_no_fallback.pptx"))
    picture = next(
        shape
        for slide in presentation.slides
        for shape in slide.shapes
        if PptxConverter()._is_picture(shape)
    )
    nonvisual = picture._element.xpath("./p:nvPicPr/p:nvPr")[0]
    etree.SubElement(
        nonvisual,
        "{http://schemas.openxmlformats.org/presentationml/2006/main}ph",
        type="pic",
        idx="1",
    )
    stream = io.BytesIO()
    presentation.save(stream)
    stream.seek(0)
    presentation = Presentation(stream)
    placeholder = next(
        shape
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER
    )
    with pytest.raises(ValueError, match="no embedded image"):
        _ = placeholder.image
    stream.seek(0)
    render = Mock(return_value="<p>SVG placeholder</p>")

    result = _ImageConverter(render).convert(stream, _INFO)

    assert "SVG placeholder" in result.markdown
    assert render.call_args.args[1].mimetype == "image/svg+xml"
    render.assert_called_once()


def test_missing_image_bytes_keep_native_placeholder(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    render = Mock(side_effect=AssertionError("unexpected image"))
    converter = _ImageConverter(render)
    monkeypatch.setattr(
        converter, "_get_image_info", Mock(return_value=(None, None, None))
    )

    result = converter.convert(_presentation(), _INFO, keep_data_uris=True)

    assert "![Alt 1 & text](Picture1.jpg)" in result.markdown
    render.assert_not_called()


@pytest.mark.parametrize("value", [False, 123, b"<p>not text</p>"])
def test_non_string_fragments_fail_explicitly(value: Any) -> None:
    with pytest.raises(TypeError, match="HTML string or None"):
        _ImageConverter(Mock(return_value=value)).convert(_presentation(), _INFO)


@pytest.mark.parametrize(
    "fragment",
    [
        "<html><body>document</body></html>",
        "<head><title>document</title></head>",
        "<body>document</body>",
        "<!DOCTYPE html><p>document</p>",
    ],
)
def test_full_documents_fail_explicitly(fragment: str) -> None:
    with pytest.raises(ValueError, match="fragment, not a document"):
        _ImageConverter(Mock(return_value=fragment)).convert(_presentation(), _INFO)


def test_hook_errors_propagate_close_stream_and_use_dispatcher_fallback() -> None:
    streams = []
    error = RuntimeError("image renderer failed")

    def render(stream, info, **kwargs):
        streams.append(stream)
        raise error

    converter = _ImageConverter(render)
    with pytest.raises(RuntimeError) as caught:
        converter.convert(_presentation(), _INFO)
    assert caught.value is error
    without_fallback = MarkItDown(enable_builtins=False)
    without_fallback.register_converter(converter)
    with pytest.raises(FileConversionException) as aggregate:
        without_fallback.convert_stream(_presentation(), stream_info=_INFO)
    assert aggregate.value.attempts is not None
    assert any(
        attempt.exc_info and attempt.exc_info[1] is error
        for attempt in aggregate.value.attempts
    )
    md = MarkItDown()
    md.register_converter(converter, priority=-1)
    assert (
        md.convert_stream(_presentation(), stream_info=_INFO).markdown
        == PptxConverter().convert(_presentation(), _INFO).markdown
    )
    assert streams and all(stream.closed for stream in streams)


@pytest.mark.parametrize("caption_result", [None, "", RuntimeError("caption failed")])
def test_caption_failure_precedes_hook_with_fresh_image_stream(
    monkeypatch: pytest.MonkeyPatch, caption_result: Any
) -> None:
    calls = []
    client = object()

    def caption(stream, info, **kwargs):
        assert stream.tell() == 0 and stream.read() == _PNG
        assert info.mimetype == "image/png" and info.extension == ".png"
        assert kwargs == {"client": client, "model": "vision", "prompt": "Describe"}
        calls.append("caption")
        if isinstance(caption_result, Exception):
            raise caption_result
        return caption_result

    def render(stream, info, **kwargs):
        assert stream.tell() == 0 and stream.read() == _PNG
        calls.append("hook")
        return "<p>recognized</p>"

    monkeypatch.setattr(_pptx_converter, "llm_caption", caption)
    result = _ImageConverter(render).convert(
        _presentation(),
        _INFO,
        llm_client=client,
        llm_model="vision",
        llm_prompt="Describe",
    )

    assert calls == ["caption", "hook"]
    assert "recognized" in result.markdown


def test_successful_caption_keeps_native_markdown_without_hook_or_html(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    caption = Mock(return_value="**Caption** [label]\nwith_under & <angle>")
    monkeypatch.setattr(_pptx_converter, "llm_caption", caption)
    render = Mock(side_effect=AssertionError("caption must take precedence"))
    converter = _ImageConverter(render)
    convert_html = Mock(side_effect=AssertionError("caption is not image HTML"))
    monkeypatch.setattr(converter._html_converter, "convert_string", convert_html)

    result = converter.convert(
        _presentation(), _INFO, llm_client=object(), llm_model="vision"
    )

    assert result.markdown == (
        "<!-- Slide number: 1 -->\n# Deck title\n\n"
        "![**Caption** label with_under & <angle> Alt 1 & text](Picture1.jpg)"
    )
    caption.assert_called_once()
    render.assert_not_called()
    convert_html.assert_not_called()


def test_missing_optional_dependencies_fail_before_hook(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    error = ImportError("python-pptx is unavailable")
    monkeypatch.setattr(
        _pptx_converter, "_dependency_exc_info", (ImportError, error, None)
    )
    render = Mock()
    with pytest.raises(MissingDependencyException) as caught:
        _ImageConverter(render).convert(io.BytesIO(b""), _INFO)
    assert caught.value.__cause__ is error
    render.assert_not_called()
