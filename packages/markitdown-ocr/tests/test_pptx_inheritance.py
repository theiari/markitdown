"""Full PPTX OCR conversions inherit core slides, image hooks, and HTML."""

import base64
import inspect
import io
from pathlib import Path
from typing import Any
from unittest.mock import Mock, create_autospec
import zipfile

from PIL import Image
from pptx import Presentation
from pptx.util import Inches
import pytest

from markitdown import FileConversionException, MarkItDown, StreamInfo
from markitdown.converters import PptxConverter
from markitdown.converters import _pptx_converter
import markitdown._markitdown as markitdown_module
from markitdown_ocr import _plugin
from markitdown_ocr._ocr_service import LLMVisionOCRService, OCRResult
from markitdown_ocr._pptx_converter_with_ocr import PptxConverterWithOCR


_INFO = StreamInfo(extension=".pptx")
_CORE_FILES = Path(__file__).parents[2] / "markitdown" / "tests" / "test_files"


def _png(color: str) -> bytes:
    stream = io.BytesIO()
    Image.new("RGB", (2, 2), color).save(stream, format="PNG")
    return stream.getvalue()


_RED = _png("red")
_BLUE = _png("blue")


def _presentation(images: tuple[bytes, ...] = (_RED,)) -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Heading"
    slide.shapes.title.top = 0
    for index in reversed(range(len(images))):
        picture = slide.shapes.add_picture(
            io.BytesIO(images[index]), 0, Inches(index + 1), width=Inches(0.5)
        )
        picture.name = f"Picture {index + 1}"
        picture._element._nvXxPr.cNvPr.set("descr", "")
    slide.notes_slide.notes_text_frame.text = "Speaker notes"
    stream = io.BytesIO()
    presentation.save(stream)
    return stream.getvalue()


def _service(text: str = "recognized") -> Mock:
    recognize = lambda stream: OCRResult(text=text)
    return Mock(extract_text=create_autospec(recognize, side_effect=recognize))


def _convert(converter: PptxConverter, data: bytes, **kwargs: Any) -> str:
    return converter.convert(io.BytesIO(data), _INFO, **kwargs).markdown


def test_pptx_ocr_is_a_thin_subclass() -> None:
    assert issubclass(PptxConverterWithOCR, PptxConverter)
    assert inspect.signature(PptxConverterWithOCR.convert) == inspect.signature(
        PptxConverter.convert
    )
    for method in (
        "accepts",
        "_is_picture",
        "_get_image_info",
        "_find_svg_blip_part",
        "_convert_picture_to_markdown",
        "_is_table",
        "_convert_table_to_markdown",
        "_convert_chart_to_markdown",
    ):
        assert getattr(PptxConverterWithOCR, method) is getattr(PptxConverter, method)


@pytest.mark.parametrize(
    ("fallback", "keep_data_uris"), [(None, False), ("", True), (" \r\n\t", False)]
)
def test_no_service_or_empty_recognition_matches_native_core(
    fallback: str | None, keep_data_uris: bool
) -> None:
    service = None if fallback is None else _service(fallback)
    data = _presentation()

    assert _convert(
        PptxConverterWithOCR(service), data, keep_data_uris=keep_data_uris
    ) == _convert(PptxConverter(), data, keep_data_uris=keep_data_uris)


def test_ocr_cache_is_document_local_and_honors_service_override() -> None:
    first = _service("first")
    second = _service("second")
    converter = PptxConverterWithOCR(first)
    data = _presentation((_RED,) * 12)

    original = _convert(converter, data)
    overridden = _convert(converter, data, ocr_service=second)
    again = _convert(converter, data)

    assert original.count("*[Image OCR]  \nfirst  \n[End OCR]*") == 12
    assert overridden.count("*[Image OCR]  \nsecond  \n[End OCR]*") == 12
    assert again == original
    assert first.extract_text.call_count == 2
    second.extract_text.assert_called_once()


def test_recognition_uses_image_identity_in_reading_order_with_native_fallback() -> (
    None
):
    calls = []

    def recognize(stream):
        assert stream.tell() == 0
        image = stream.read()
        calls.append(image)
        return OCRResult(text="blue" if image == _BLUE else "")

    service = Mock(extract_text=create_autospec(recognize, side_effect=recognize))
    result = _convert(
        PptxConverterWithOCR(service), _presentation((_BLUE, _RED, _BLUE))
    )

    assert calls == [_BLUE, _RED]
    assert result.count("*[Image OCR]") == 2
    assert result.count("![Picture 2](Picture2.jpg)") == 1
    assert result.index("blue") < result.index("![Picture 2]") < result.rindex("blue")
    assert result.endswith("### Notes:\nSpeaker notes")


def test_ocr_text_is_escaped_with_normalized_html_breaks_and_html_options(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    converter = PptxConverterWithOCR(_service("A_B *literal*\r\n<tag> & value\rfinal"))
    convert_html = Mock(wraps=converter._html_converter.convert_string)
    monkeypatch.setattr(converter._html_converter, "convert_string", convert_html)
    data = _presentation()

    result = _convert(converter, data)

    assert convert_html.call_args.args == (
        "<p><em>[Image OCR]<br/>A_B *literal*<br/>&lt;tag&gt; &amp; value"
        "<br/>final<br/>[End OCR]</em></p>",
    )
    assert (
        "*[Image OCR]  \nA\\_B \\*literal\\*  \n" "<tag> & value  \nfinal  \n[End OCR]*"
    ) in result
    assert "*literal*" in _convert(
        converter, data, escape_asterisks=False, escape_underscores=False
    )
    assert "A_B" in _convert(converter, data, escape_underscores=False)


def test_reported_errors_warn_once_per_image_and_keep_native_output() -> None:
    service = Mock(
        extract_text=Mock(
            return_value=OCRResult(text="discard", error="quota exceeded")
        )
    )
    data = _presentation((_RED, _RED))

    with pytest.warns(
        RuntimeWarning, match="PPTX image OCR failed: quota exceeded"
    ) as seen:
        result = _convert(PptxConverterWithOCR(service), data, keep_data_uris=True)

    assert len(seen) == 1
    assert result == _convert(PptxConverter(), data, keep_data_uris=True)
    service.extract_text.assert_called_once()


def test_thrown_errors_propagate_and_use_normal_dispatcher_fallback() -> None:
    error = RuntimeError("custom OCR failed")
    converter = PptxConverterWithOCR(Mock(extract_text=Mock(side_effect=error)))
    data = _presentation()
    with pytest.raises(RuntimeError) as caught:
        _convert(converter, data)
    assert caught.value is error
    md = MarkItDown()
    md.register_converter(converter, priority=-1)
    assert md.convert_stream(io.BytesIO(data), stream_info=_INFO).markdown == _convert(
        PptxConverter(), data
    )
    without_fallback = MarkItDown(enable_builtins=False)
    without_fallback.register_converter(converter)
    with pytest.raises(FileConversionException) as aggregate:
        without_fallback.convert_stream(io.BytesIO(data), stream_info=_INFO)
    assert aggregate.value.attempts is not None
    assert any(
        attempt.exc_info and attempt.exc_info[1] is error
        for attempt in aggregate.value.attempts
    )


def test_svg_only_images_use_inherited_resolution_and_ocr() -> None:
    seen = []

    def recognize(stream):
        seen.append(stream.read())
        return OCRResult(text="SVG text")

    data = (_CORE_FILES / "test_svg_no_fallback.pptx").read_bytes()
    result = _convert(
        PptxConverterWithOCR(
            Mock(extract_text=create_autospec(recognize, side_effect=recognize))
        ),
        data,
    )

    assert len(seen) == 1 and b"<svg" in seen[0]
    assert "*[Image OCR]  \nSVG text  \n[End OCR]*" in result
    assert "data:image" not in result


def test_svg_metadata_reaches_bundled_vision_request() -> None:
    data = (_CORE_FILES / "test_svg_no_fallback.pptx").read_bytes()
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        svg_parts = [name for name in archive.namelist() if name.endswith(".svg")]
        assert len(svg_parts) == 1
        svg = archive.read(svg_parts[0])
    client = Mock()
    client.chat.completions.create.return_value.choices = [
        Mock(message=Mock(content="SVG text"))
    ]
    service = LLMVisionOCRService(client, "vision-model")

    result = _convert(PptxConverterWithOCR(service), data)

    assert "*[Image OCR]  \nSVG text  \n[End OCR]*" in result
    client.chat.completions.create.assert_called_once()
    request = client.chat.completions.create.call_args.kwargs
    assert request["model"] == "vision-model"
    assert request["messages"][0]["content"][1]["image_url"]["url"] == (
        "data:image/svg+xml;base64," + base64.b64encode(svg).decode("ascii")
    )


@pytest.mark.parametrize("caption_succeeds", [False, True])
def test_plugin_full_trip_with_mocked_model_only(
    monkeypatch: pytest.MonkeyPatch, caption_succeeds: bool
) -> None:
    entry_point = Mock()
    entry_point.load.return_value = _plugin
    entry_points = Mock(return_value=[entry_point])
    monkeypatch.setattr(markitdown_module, "entry_points", entry_points)
    monkeypatch.setattr(markitdown_module, "_plugins", None)
    client = Mock()
    response = Mock(choices=[Mock(message=Mock(content="Recognized_text"))])
    client.chat.completions.create.side_effect = (
        [response]
        if caption_succeeds
        else [Mock(choices=[Mock(message=Mock(content=None))]), response]
    )
    md = MarkItDown(
        enable_plugins=True,
        llm_client=client,
        llm_model="vision-model",
        llm_prompt="Read the text",
    )
    registered = [
        registration
        for registration in md._converters
        if isinstance(registration.converter, PptxConverterWithOCR)
    ]
    assert len(registered) == 1 and registered[0].priority == -1

    result = md.convert_stream(io.BytesIO(_presentation()), stream_info=_INFO)

    expected_image = (
        "![Recognized_text](Picture1.jpg)"
        if caption_succeeds
        else "*[Image OCR]\nRecognized\\_text\n[End OCR]*"
    )
    assert result.markdown == (
        "<!-- Slide number: 1 -->\n# Heading\n\n"
        + expected_image
        + "\n\n### Notes:\nSpeaker notes"
    )
    entry_points.assert_called_once_with(group="markitdown.plugin")
    assert client.chat.completions.create.call_count == (1 if caption_succeeds else 2)
    for call in client.chat.completions.create.call_args_list:
        request = call.kwargs
        assert request["model"] == "vision-model"
        content = request["messages"][0]["content"]
        assert content[0]["text"] == "Read the text"
        assert content[1]["image_url"]["url"] == (
            "data:image/png;base64," + base64.b64encode(_RED).decode("ascii")
        )


def test_native_caption_precedes_ocr_and_does_not_escape_existing_markdown(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    caption = Mock(return_value="**Caption** with_under & <literal>")
    monkeypatch.setattr(_pptx_converter, "llm_caption", caption)
    service = _service()
    data = _presentation()
    options = {"llm_client": object(), "llm_model": "vision"}

    result = _convert(PptxConverterWithOCR(service), data, **options)

    assert result == _convert(PptxConverter(), data, **options)
    assert "![**Caption** with_under & <literal>](Picture1.jpg)" in result
    service.extract_text.assert_not_called()
    assert caption.call_count == 2


def test_older_core_fails_explicitly(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.delattr(_pptx_converter.PptxConverter, "_image_to_html")

    with pytest.raises(RuntimeError, match=r"markitdown>=0\.1\.8b3") as caught:
        PptxConverterWithOCR()
    assert "pip install --upgrade 'markitdown>=0.1.8b3'" in str(caught.value)
