"""Full DOCX OCR trips through the inherited core image/HTML pipeline."""

import base64
import inspect
import io
from typing import Any
from unittest.mock import Mock, create_autospec

from bs4 import BeautifulSoup
from docx import Document
from PIL import Image
import pytest

from markitdown import FileConversionException, MarkItDown, StreamInfo
from markitdown.converters import DocxConverter
from markitdown.converters import _docx_converter
import markitdown._markitdown as markitdown_module
from markitdown_ocr import _plugin
from markitdown_ocr._docx_converter_with_ocr import DocxConverterWithOCR
from markitdown_ocr._ocr_service import OCRResult


_INFO = StreamInfo(extension=".docx")


def _png(color: str) -> bytes:
    output = io.BytesIO()
    Image.new("RGB", (2, 2), color).save(output, format="PNG")
    return output.getvalue()


_RED = _png("red")
_BLUE = _png("blue")


def _document(
    image_data: tuple[bytes, ...] = (_RED,),
    *,
    in_table: bool = False,
    inline: bool = False,
) -> bytes:
    document = Document()
    document.add_heading("Heading", level=1)
    for data in image_data:
        if in_table:
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Item"
            table.cell(0, 1).text = "Details"
            table.cell(1, 0).text = "A"
            paragraph = table.cell(1, 1).paragraphs[0]
        else:
            paragraph = document.add_paragraph()
        if inline:
            paragraph.add_run("Before ")
        paragraph.add_run().add_picture(io.BytesIO(data))
        if inline:
            paragraph.add_run(" After")
    document.add_paragraph("Native content").runs[0].bold = True
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _service(text: str = "recognized") -> Mock:
    # A one-argument service must remain supported, without injected keywords.
    recognize = lambda stream: OCRResult(text=text)
    return Mock(extract_text=create_autospec(recognize, side_effect=recognize))


def _convert(converter: DocxConverter, data: bytes, **kwargs: Any) -> str:
    return converter.convert(io.BytesIO(data), _INFO, **kwargs).markdown


def test_docx_ocr_is_a_thin_subclass() -> None:
    assert issubclass(DocxConverterWithOCR, DocxConverter)
    assert DocxConverterWithOCR.accepts is DocxConverter.accepts
    assert inspect.signature(DocxConverterWithOCR.convert) == inspect.signature(
        DocxConverter.convert
    )
    assert not hasattr(DocxConverterWithOCR, "_inject_placeholders")
    assert not hasattr(DocxConverterWithOCR, "_extract_and_ocr_images")


def test_plugin_registration_full_trip(monkeypatch: pytest.MonkeyPatch) -> None:
    entry_point = Mock()
    entry_point.load.return_value = _plugin
    entry_points = Mock(return_value=[entry_point])
    monkeypatch.setattr(markitdown_module, "entry_points", entry_points)
    monkeypatch.setattr(markitdown_module, "_plugins", None)
    client = Mock()
    client.chat.completions.create.return_value.choices = [
        Mock(message=Mock(content="Recognized_text"))
    ]
    md = MarkItDown(
        enable_plugins=True,
        llm_client=client,
        llm_model="vision-model",
        llm_prompt="Read the text",
    )
    registered = [
        registration
        for registration in md._converters
        if isinstance(registration.converter, DocxConverterWithOCR)
    ]
    assert len(registered) == 1 and registered[0].priority == -1

    result = md.convert_stream(io.BytesIO(_document(inline=True)), stream_info=_INFO)

    assert result.markdown == (
        "# Heading\n\nBefore\n\n"
        "*[Image OCR]\nRecognized\\_text\n[End OCR]*\n\n"
        "After\n\n**Native content**"
    )
    entry_points.assert_called_once_with(group="markitdown.plugin")
    client.chat.completions.create.assert_called_once()
    request = client.chat.completions.create.call_args.kwargs
    assert request["model"] == "vision-model"
    content = request["messages"][0]["content"]
    assert content[0]["text"] == "Read the text"
    assert content[1]["image_url"]["url"] == (
        "data:image/png;base64," + base64.b64encode(_RED).decode("ascii")
    )


@pytest.mark.parametrize("keep_data_uris", [False, True])
@pytest.mark.parametrize("fallback", [None, "", " \n\t"])
def test_no_service_or_empty_text_matches_core(
    keep_data_uris: bool, fallback: str | None
) -> None:
    service = None if fallback is None else _service(fallback)
    data = _document(inline=True)
    expected = _convert(DocxConverter(), data, keep_data_uris=keep_data_uris)

    actual = _convert(
        DocxConverterWithOCR(ocr_service=service),
        data,
        keep_data_uris=keep_data_uris,
    )

    assert actual == expected


def test_recognition_cache_is_local_and_service_override_is_preserved() -> None:
    first = _service("first")
    second = _service("second")
    converter = DocxConverterWithOCR(ocr_service=first)
    data = _document((_RED,) * 12)

    default = _convert(converter, data)
    overridden = _convert(converter, data, ocr_service=second)
    again = _convert(converter, data)

    assert default.count("*[Image OCR]  \nfirst  \n[End OCR]*") == 12
    assert overridden.count("*[Image OCR]  \nsecond  \n[End OCR]*") == 12
    assert again == default
    assert first.extract_text.call_count == 2
    second.extract_text.assert_called_once()


def test_image_identity_not_relationship_order_and_failed_images_stay_native() -> None:
    document = Document(io.BytesIO(_document((_RED, _BLUE, _RED))))
    red, blue = document.paragraphs[1:3]
    red._p.addprevious(blue._p)
    stream = io.BytesIO()
    document.save(stream)
    calls = []

    def recognize(image_stream):
        image = image_stream.read()
        calls.append(image)
        return OCRResult(text="blue" if image == _BLUE else "")

    service = Mock(extract_text=create_autospec(recognize, side_effect=recognize))
    result = _convert(DocxConverterWithOCR(service), stream.getvalue())

    assert calls == [_BLUE, _RED]
    assert result.count("*[Image OCR]") == 1
    assert result.count("data:image/png;base64...") == 2
    assert result.index("blue") < result.index("data:image/png;base64...")


def test_ocr_text_is_escaped_and_line_endings_are_preserved_as_html_breaks(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    service = _service("A_B *literal*\r\n<tag> & value\rfinal")
    converter = DocxConverterWithOCR(service)
    convert_html = Mock(wraps=converter._html_converter.convert_string)
    monkeypatch.setattr(converter._html_converter, "convert_string", convert_html)

    result = _convert(converter, _document())

    assert (
        "<p><em>[Image OCR]<br/>A_B *literal*<br/>&lt;tag&gt; &amp; value"
        "<br/>final<br/>[End OCR]</em></p>"
    ) in convert_html.call_args.args[0]
    assert (
        "*[Image OCR]  \nA\\_B \\*literal\\*  \n" "<tag> & value  \nfinal  \n[End OCR]*"
    ) in result
    assert "data-markitdown-image-" not in convert_html.call_args.args[0]


def test_ocr_inside_a_table_reaches_html_conversion_inside_the_cell(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    converter = DocxConverterWithOCR(_service("Serial: 12345\nStatus: active"))
    convert_html = Mock(wraps=converter._html_converter.convert_string)
    monkeypatch.setattr(converter._html_converter, "convert_string", convert_html)

    result = _convert(converter, _document(in_table=True))

    soup = BeautifulSoup(convert_html.call_args.args[0], "html.parser")
    assert len(soup.find_all("td")) == 4
    assert str(soup.find_all("td")[-1]) == (
        "<td><p><em>[Image OCR]<br/>Serial: 12345<br/>"
        "Status: active<br/>[End OCR]</em></p></td>"
    )
    assert not soup.select("p p")
    assert "| A | *[Image OCR] Serial: 12345 Status: active [End OCR]* |" in result
    assert "**Native content**" in result


def test_style_maps_and_html_options_are_inherited() -> None:
    document = Document(io.BytesIO(_document()))
    document.add_paragraph().add_run("underlined").underline = True
    stream = io.BytesIO()
    document.save(stream)

    result = _convert(
        DocxConverterWithOCR(_service("OCR_text")),
        stream.getvalue(),
        style_map="u => strong",
        escape_underscores=False,
        heading_style="underlined",
    )

    assert result.startswith("Heading\n=======")
    assert "OCR_text" in result
    assert "**underlined**" in result


def test_reported_ocr_error_warns_and_keeps_native_images() -> None:
    service = Mock(
        extract_text=Mock(return_value=OCRResult(text="", error="quota exceeded"))
    )
    data = _document((_RED, _RED))

    with pytest.warns(RuntimeWarning, match="quota exceeded") as warnings:
        result = _convert(DocxConverterWithOCR(service), data)

    assert len(warnings) == 1
    assert result == _convert(DocxConverter(), data)
    service.extract_text.assert_called_once()


def test_raised_service_error_uses_normal_dispatcher_fallback() -> None:
    error = RuntimeError("custom service failed")
    service = Mock(extract_text=Mock(side_effect=error))
    converter = DocxConverterWithOCR(service)
    data = _document()
    with pytest.raises(RuntimeError) as caught:
        _convert(converter, data)
    assert caught.value is error

    md = MarkItDown()
    md.register_converter(converter, priority=-1)
    assert md.convert_stream(io.BytesIO(data), stream_info=_INFO).markdown == _convert(
        DocxConverter(), data
    )
    without_fallback = MarkItDown(enable_builtins=False)
    without_fallback.register_converter(converter, priority=-1)
    with pytest.raises(FileConversionException) as aggregate:
        without_fallback.convert_stream(io.BytesIO(data), stream_info=_INFO)
    assert aggregate.value.attempts is not None
    assert any(
        attempt.exc_info and attempt.exc_info[1] is error
        for attempt in aggregate.value.attempts
    )


def test_older_core_is_rejected_instead_of_silently_skipping_ocr(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.delattr(_docx_converter.DocxConverter, "_image_to_html")

    with pytest.raises(RuntimeError, match=r"markitdown>=0\.1\.8b3") as caught:
        DocxConverterWithOCR()
    assert "pip install --upgrade 'markitdown>=0.1.8b3'" in str(caught.value)
