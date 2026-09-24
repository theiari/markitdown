"""XLSX OCR trips through native tables, repair, extraction and HTML conversion."""

import base64
import inspect
import io
from typing import Any
from unittest.mock import Mock, create_autospec
import zipfile

import openpyxl
from openpyxl.drawing.image import Image as SheetImage
from openpyxl.drawing.spreadsheet_drawing import (
    AbsoluteAnchor,
    AnchorMarker,
    OneCellAnchor,
    TwoCellAnchor,
)
from PIL import Image
import pytest

from markitdown import FileConversionException, MarkItDown, StreamInfo
from markitdown.converters import XlsxConverter
from markitdown.converters import _xlsx_converter
import markitdown._markitdown as markitdown_module
from markitdown_ocr import _plugin
from markitdown_ocr._ocr_service import OCRResult
from markitdown_ocr._xlsx_converter_with_ocr import XlsxConverterWithOCR


_INFO = StreamInfo(extension=".xlsx")


def _png(color: str) -> bytes:
    stream = io.BytesIO()
    Image.new("RGB", (2, 2), color).save(stream, "PNG")
    return stream.getvalue()


_RED, _BLUE = _png("red"), _png("blue")


def _workbook(images: tuple[bytes, ...] = (_RED,)) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Cells"
    sheet.append(["Header_one", None])
    sheet.append(["<native> & *value*", 12])
    for index, data in enumerate(images):
        sheet.add_image(SheetImage(io.BytesIO(data)), f"AA{index + 3}")
    last = workbook.create_sheet("Other")
    last.append(["Last"])
    last.append(["Final"])
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def _service(text: str = "recognized") -> Mock:
    recognize = lambda stream: OCRResult(text=text)
    return Mock(extract_text=create_autospec(recognize, side_effect=recognize))


def _convert(converter: XlsxConverter, data: bytes, **kwargs: Any) -> str:
    return converter.convert(io.BytesIO(data), _INFO, **kwargs).markdown


def test_ocr_is_a_thin_subclass_with_native_acceptance_and_signature() -> None:
    assert issubclass(XlsxConverterWithOCR, XlsxConverter)
    assert XlsxConverterWithOCR.accepts is XlsxConverter.accepts
    assert inspect.signature(XlsxConverterWithOCR.convert) == inspect.signature(
        XlsxConverter.convert
    )
    assert not hasattr(XlsxConverterWithOCR, "_convert_standard")
    assert not hasattr(XlsxConverterWithOCR, "_convert_with_ocr")
    assert not hasattr(XlsxConverterWithOCR, "_extract_and_ocr_sheet_images")
    assert not XlsxConverterWithOCR().accepts(
        io.BytesIO(), StreamInfo(extension=".xls")
    )


def test_plugin_full_trip_uses_native_cells_and_original_image(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    entry_point = Mock()
    entry_point.load.return_value = _plugin
    monkeypatch.setattr(
        markitdown_module, "entry_points", Mock(return_value=[entry_point])
    )
    monkeypatch.setattr(markitdown_module, "_plugins", None)
    client = Mock()
    client.chat.completions.create.return_value.choices = [
        Mock(message=Mock(content="Recognized_text"))
    ]
    md = MarkItDown(
        enable_plugins=True,
        llm_client=client,
        llm_model="vision-model",
        llm_prompt="Read the text",
    )
    registered = [
        registration
        for registration in md._converters
        if isinstance(registration.converter, XlsxConverterWithOCR)
    ]
    assert len(registered) == 1 and registered[0].priority == -1
    result = md.convert_stream(io.BytesIO(_workbook((_RED, _RED))), stream_info=_INFO)

    assert result.markdown == (
        "## Cells\n"
        "| Header\\_one | Unnamed: 1 |\n"
        "| --- | --- |\n"
        "| <native> & \\*value\\* | 12 |\n\n"
        "### Images in this sheet:\n\n"
        "*[Image OCR]\nRecognized\\_text\n[End OCR]*\n\n"
        "*[Image OCR]\nRecognized\\_text\n[End OCR]*\n\n"
        "## Other\n| Last |\n| --- |\n| Final |"
    )
    client.chat.completions.create.assert_called_once()
    request = client.chat.completions.create.call_args.kwargs
    assert request["model"] == "vision-model"
    content = request["messages"][0]["content"]
    assert content[0]["text"] == "Read the text"
    assert content[1]["image_url"]["url"] == (
        "data:image/png;base64," + base64.b64encode(_RED).decode("ascii")
    )


@pytest.mark.parametrize("text", [None, "", " \n\t"])
def test_no_service_or_blank_recognition_is_exactly_native(text: str | None) -> None:
    service = None if text is None else _service(text)
    data = _workbook()
    options = {"escape_underscores": False, "heading_style": "underlined"}
    assert _convert(XlsxConverterWithOCR(service), data, **options) == (
        _convert(XlsxConverter(), data, **options)
    )


def test_no_images_does_not_call_service() -> None:
    service = _service()
    data = _workbook(())
    assert _convert(XlsxConverterWithOCR(service), data) == _convert(
        XlsxConverter(), data
    )
    service.extract_text.assert_not_called()


def test_cache_is_document_local_and_per_call_service_overrides_work() -> None:
    first = _service("first")
    second = _service("second")
    converter = XlsxConverterWithOCR(first)
    data = _workbook((_RED,) * 4)
    default = _convert(converter, data)
    overridden = _convert(converter, data, ocr_service=second)
    again = _convert(converter, data)

    assert default.count("*[Image OCR]  \nfirst  \n[End OCR]*") == 4
    assert overridden.count("*[Image OCR]  \nsecond  \n[End OCR]*") == 4
    assert again == default
    assert first.extract_text.call_count == 2
    second.extract_text.assert_called_once()


def test_recognition_identity_and_blank_images_preserve_anchor_order() -> None:
    calls = []

    def recognize(stream):
        data = stream.read()
        calls.append(data)
        return OCRResult(text="blue" if data == _BLUE else "")

    converter = XlsxConverterWithOCR(
        Mock(extract_text=create_autospec(recognize, side_effect=recognize))
    )
    result = _convert(converter, _workbook((_RED, _BLUE, _RED, _BLUE)))

    assert calls == [_RED, _BLUE]
    assert result.count("[Image OCR]") == 2
    assert "Image at " not in result
    assert result.count("*[Image OCR]  \nblue  \n[End OCR]*") == 2


def test_mixed_anchor_recognition_keeps_legacy_openpyxl_order() -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Native"])
    sheet.append(["Cells"])
    images = [
        ("one first", _RED, OneCellAnchor(_from=AnchorMarker(col=1, row=3))),
        ("absolute", _BLUE, AbsoluteAnchor()),
        (
            "two",
            _png("green"),
            TwoCellAnchor(
                _from=AnchorMarker(col=0, row=1), to=AnchorMarker(col=4, row=4)
            ),
        ),
        ("one second", _png("yellow"), OneCellAnchor(_from=AnchorMarker(col=3, row=7))),
    ]
    labels = {}
    for label, data, anchor in images:
        labels[data] = label
        image = SheetImage(io.BytesIO(data))
        image.anchor = anchor
        sheet.add_image(image)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    data = output.getvalue()
    legacy = openpyxl.load_workbook(io.BytesIO(data))
    try:
        legacy_order = [labels[image._data()] for image in legacy.active._images]
    finally:
        legacy.close()
    assert legacy_order == ["absolute", "one first", "one second", "two"]
    calls = []

    def recognize(stream):
        label = labels[stream.read()]
        calls.append(label)
        return OCRResult(text=label)

    result = _convert(
        XlsxConverterWithOCR(
            Mock(extract_text=create_autospec(recognize, side_effect=recognize))
        ),
        data,
    )
    assert calls == legacy_order
    blocks = [f"*[Image OCR]  \n{label}  \n[End OCR]*" for label in legacy_order]
    assert "\n\n".join(blocks) in result


def test_ocr_html_is_escaped_and_existing_html_options_are_forwarded(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    converter = XlsxConverterWithOCR(_service("A_B *literal*\r\n<tag> & value\rfinal"))
    convert_html = Mock(wraps=converter._html_converter.convert_string)
    monkeypatch.setattr(converter._html_converter, "convert_string", convert_html)

    result = _convert(
        converter, _workbook(), escape_underscores=False, heading_style="underlined"
    )

    assert (
        "<p><em>[Image OCR]<br/>A_B *literal*<br/>&lt;tag&gt; &amp; value"
        "<br/>final<br/>[End OCR]</em></p>"
    ) in convert_html.call_args_list[1].args[0]
    assert (
        "*[Image OCR]  \nA_B \\*literal\\*  \n<tag> & value  \nfinal  \n[End OCR]*"
        in result
    )
    assert "| Header_one | Unnamed: 1 |" in result
    assert all(
        call.kwargs["escape_underscores"] is False
        for call in convert_html.call_args_list
    )


def test_inherited_repairs_and_native_table_fixes_reach_ocr(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    data = _workbook((_RED, _RED))
    stream = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source:
        with zipfile.ZipFile(stream, "w") as target:
            for entry in source.infolist():
                content = source.read(entry)
                if entry.filename == "xl/worksheets/sheet1.xml":
                    content = content.replace(
                        b"<sheetView ", b'<sheetView showZeroes="0" ', 1
                    )
                target.writestr(entry, content)
    service = _service()
    repair = Mock(wraps=_xlsx_converter._repair_sheetview_show_zeroes)
    monkeypatch.setattr(_xlsx_converter, "_repair_sheetview_show_zeroes", repair)
    native_read = _xlsx_converter.pd.read_excel
    calls = []

    def fixed_read(*args: Any, **kwargs: Any):
        calls.append(kwargs)
        sheets = native_read(*args, **kwargs)
        sheets["Cells"].iloc[0, 0] = "Future shared fix"
        return sheets

    monkeypatch.setattr(_xlsx_converter.pd, "read_excel", fixed_read)
    result = _convert(XlsxConverterWithOCR(service), stream.getvalue())
    assert "Future shared fix" in result and "<native>" not in result
    assert result.count("[Image OCR]") == 2
    repair.assert_called_once()
    assert calls == [{"sheet_name": None, "engine": "openpyxl"}] * 2
    service.extract_text.assert_called_once()


def test_reported_ocr_error_warns_once_and_keeps_native_output() -> None:
    service = Mock(
        extract_text=Mock(
            return_value=OCRResult(text="ignored", error="quota exceeded")
        )
    )
    data = _workbook((_RED, _RED))
    with pytest.warns(RuntimeWarning, match="quota exceeded") as warnings:
        result = _convert(XlsxConverterWithOCR(service), data)
    assert len(warnings) == 1
    assert result == _convert(XlsxConverter(), data)
    service.extract_text.assert_called_once()


def test_thrown_ocr_errors_use_standard_dispatcher_fallback() -> None:
    error = RuntimeError("custom service failed")
    converter = XlsxConverterWithOCR(Mock(extract_text=Mock(side_effect=error)))
    data = _workbook()
    with pytest.raises(RuntimeError) as caught:
        _convert(converter, data)
    assert caught.value is error

    md = MarkItDown()
    md.register_converter(converter, priority=-1)
    assert md.convert_stream(io.BytesIO(data), stream_info=_INFO).markdown == _convert(
        XlsxConverter(), data
    )
    no_fallback = MarkItDown(enable_builtins=False)
    no_fallback.register_converter(converter, priority=-1)
    with pytest.raises(FileConversionException) as aggregate:
        no_fallback.convert_stream(io.BytesIO(data), stream_info=_INFO)
    assert aggregate.value.attempts is not None
    assert any(
        attempt.exc_info and attempt.exc_info[1] is error
        for attempt in aggregate.value.attempts
    )


def test_older_core_fails_with_actionable_upgrade_message(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.delattr(_xlsx_converter.XlsxConverter, "_image_to_html")
    with pytest.raises(RuntimeError, match=r"markitdown>=0\.1\.8b3") as caught:
        XlsxConverterWithOCR()
    assert "pip install --upgrade 'markitdown>=0.1.8b3'" in str(caught.value)
